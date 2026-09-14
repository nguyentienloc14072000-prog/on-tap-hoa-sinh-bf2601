from pathlib import Path
import json
from pypdf import PdfReader
import logging
logging.getLogger('pypdf').setLevel(logging.ERROR)
from docx import Document
from zipfile import ZipFile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parents[1] / 'tmp' / 'sources'
OUT.mkdir(parents=True, exist_ok=True)
catalog = []
for folder in ['Slide 2024']:
    for path in sorted((ROOT / folder).iterdir()):
        pages = []
        if path.suffix.lower() == '.pdf':
            pages = [p.extract_text() or '' for p in PdfReader(path).pages]
        elif path.suffix.lower() == '.docx':
            doc = Document(path)
            pages = ['\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)]
        elif path.suffix.lower() == '.pptx':
            with ZipFile(path) as z:
                names = sorted([n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')], key=lambda x:int(x.split('slide')[-1].split('.')[0]))
                pages = ['\n'.join(e.text or '' for e in ET.fromstring(z.read(n)).iter() if e.tag.endswith('}t')) for n in names]
        else:
            continue
        item = {'id': f's{len(catalog)+1:02}', 'file': str(path.relative_to(ROOT)), 'pages': pages}
        catalog.append(item)
        (OUT / (item['id']+'.txt')).write_text('\n\n'.join(f'=== TRANG {i+1} ===\n{p}' for i,p in enumerate(pages)), encoding='utf-8')
        print(item['id'], str(path.relative_to(ROOT)), 'pages:',len(pages),'chars:',sum(map(len,pages)), flush=True)
(OUT / 'catalog.json').write_text(json.dumps(catalog,ensure_ascii=False,indent=2),encoding='utf-8')
